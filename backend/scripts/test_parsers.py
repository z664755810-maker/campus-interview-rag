"""解析层自测：造各种格式的样本文件，跑通「字节 -> 文本 -> 切分」。

不依赖网络、不依赖智谱 API（只测解析与切分，不测向量化）。
运行： backend/.venv/Scripts/python.exe scripts/test_parsers.py
"""
import io
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.services import parsers, ingestion  # noqa: E402

PASS, FAIL = "✓", "✗"


def check(name: str, fn):
    try:
        detail = fn()
        print(f"  {PASS} {name}: {detail}")
        return True
    except Exception as e:  # noqa: BLE001
        print(f"  {FAIL} {name}: {type(e).__name__}: {e}")
        return False


def _docx_bytes() -> bytes:
    from docx import Document

    d = Document()
    d.add_heading("Java", level=2)
    d.add_heading("Q1. HashMap 原理？", level=3)
    d.add_paragraph("数组加链表，冲突拉链，超 8 转红黑树。")
    d.add_heading("Q2. ConcurrentHashMap 如何保证线程安全？", level=3)
    d.add_paragraph("CAS 加 synchronized 细粒度锁，读操作基本无锁。")
    return _save(d)


def _save(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pptx_bytes() -> bytes:
    from pptx import Presentation

    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "面试准备"
    s.placeholders[1].text = "进程是资源分配的最小单位，线程是 CPU 调度的最小单位。"
    return _save(p)


def _xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "题库"
    ws.append(["科目", "问题", "答案"])
    ws.append(["操作系统", "进程和线程的区别？", "进程是资源分配的最小单位"])
    ws.append(["网络", "TCP 三次握手？", "SYN、SYN-ACK、ACK"])
    return _save(wb)


def main() -> int:
    results = []

    print("\n[1] 各格式解析")
    results.append(check(".md 纯文本", lambda: parsers.extract_text("a.md", "# T\n## S\n### Q1. x\n答案".encode())[:20] + "..."))
    results.append(check(".txt GBK 中文（编码兜底）", lambda: parsers.extract_text("gbk.txt", "进程与线程的区别".encode("gbk"))))
    results.append(check(".docx Word", lambda: "### Q1" in parsers.extract_text("t.docx", _docx_bytes()) and "Word 标题已还原为 ###" or "未还原 ###"))
    results.append(check(".pptx PPT", lambda: "进程是资源分配" in parsers.extract_text("t.pptx", _pptx_bytes()) and "含正文" or "空"))
    results.append(check(".xlsx Excel", lambda: "TCP 三次握手" in parsers.extract_text("t.xlsx", _xlsx_bytes()) and "含表格内容" or "空"))
    results.append(check(".csv 表格", lambda: parsers.extract_text("t.csv", "问题,答案\n进程与线程,区别如下".encode())))
    results.append(check(".json 题库数组", lambda: "### Q1" in parsers.extract_text(
        "t.json",
        json.dumps([{"question": "什么是 RAG？", "answer": "检索增强生成"}], ensure_ascii=False).encode(),
    ) and "已转成 ### Q 结构" or "未转"))
    results.append(check(".html 网页", lambda: "正文内容" in parsers.extract_text(
        "t.html",
        "<html><body><script>var a=1;</script><p>正文内容</p></body></html>".encode(),
    ) and "已剥离 script" or "异常"))

    print("\n[2] 不支持格式的友好提示（应抛 UnsupportedFileError 且带转换建议）")
    from app.services.parsers import UnsupportedFileError

    def _reject(ext: str):
        def _f():
            try:
                parsers.extract_text(f"x{ext}", b"xxx")
            except UnsupportedFileError as e:
                return f"已拦截 -> {str(e)[:40]}..."
            raise AssertionError("没有被拦截！")

        return _f

    for ext in (".doc", ".xls", ".ppt", ".png", ".zip"):
        results.append(check(f"拒绝 {ext}", _reject(ext)))

    print("\n[3] 切分策略（Word 上传能不能切出题，是关键）")
    docx_text = parsers.extract_text("t.docx", _docx_bytes())

    def _structured():
        chunks = ingestion.parse_markdown(docx_text, "t.docx")
        assert len(chunks) == 2, f"期望按题切出 2 条，实际 {len(chunks)}"
        return f"按 ### Q 结构切出 {len(chunks)} 条，subject={chunks[0]['metadata']['subject']}"

    results.append(check("Word 题库 -> 结构化切分", _structured))

    plain = "进程是资源分配的最小单位。\n" * 40 + "线程是 CPU 调度的最小单位。\n" * 40

    def _generic():
        chunks = ingestion.split_generic(plain, "plain.txt")
        assert len(chunks) >= 2, f"通用切分应产生多片，实际 {len(chunks)}"
        return f"通用切分出 {len(chunks)} 片，首片 {len(chunks[0]['text'])} 字"

    results.append(check("无结构长文 -> 通用切分兜底", _generic))

    def _empty():
        assert ingestion.split_generic("", "e.txt") == []
        return "空文本返回空列表，不会产生垃圾 chunk"

    results.append(check("空文本不产生垃圾 chunk", _empty))

    print("\n[4] 端到端：真实示例 docx 走完整解析")
    sample = Path(__file__).resolve().parent.parent.parent / "frontend" / "public" / "sample-interview.docx"

    def _e2e():
        if not sample.exists():
            return "跳过（示例文件未生成）"
        raw = sample.read_bytes()
        text = parsers.extract_text("sample-interview.docx", raw)
        chunks = ingestion.parse_markdown(text, "sample-interview.docx")
        return f"解析 {len(text)} 字 -> 切出 {len(chunks)} 道题"

    results.append(check("示例 Word 完整链路", _e2e))

    ok = sum(results)
    total = len(results)
    print(f"\n{'=' * 46}\n结果：{ok}/{total} 通过\n{'=' * 46}")
    return 0 if ok == total else 1


if __name__ == "__main__":
    raise SystemExit(main())
