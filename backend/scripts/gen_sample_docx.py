"""把内置 sample_interview.md 转成一份「真实 Word 题库」样例，供前端下载。

为什么要生成 .docx 而不是只留 .md：
  演示「支持 Word 上传」最有力的证据，就是让用户真的拖一份 Word 进去。
  这份 docx 保留了 Heading 1/2/3 层级，正好验证解析器的
  「Word 标题样式 -> Markdown # 层级 -> 按题切分」这条链路。
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent  # backend/
SRC = ROOT / "data" / "sample_interview.md"
OUT_DIR = Path(__file__).resolve().parent.parent.parent / "frontend" / "public"
OUT_DOCX = OUT_DIR / "sample-interview.docx"
OUT_MD = OUT_DIR / "sample-interview.md"


def build() -> None:
    if not SRC.exists():
        raise SystemExit(f"找不到源题库：{SRC}")

    text = SRC.read_text(encoding="utf-8")
    doc = Document()

    # 中文字体：设置东亚字体，避免 Word 里中文显示为方框
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(10.5)
    if style.element.rPr is not None and style.element.rPr.rFonts is not None:
        style.element.rPr.rFonts.set(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia",
            "Microsoft YaHei",
        )

    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        # 跳过 md 里的引用说明块，Word 里不需要
        if stripped.startswith(">"):
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        else:
            doc.add_paragraph(stripped)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    # md 示例也拷一份到 public，方便直接下载
    OUT_MD.write_text(text, encoding="utf-8")

    print(f"[ok] {OUT_DOCX}  ({OUT_DOCX.stat().st_size / 1024:.1f} KB)")
    print(f"[ok] {OUT_MD}  ({OUT_MD.stat().st_size / 1024:.1f} KB)")


if __name__ == "__main__":
    build()
